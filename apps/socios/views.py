from django.contrib import messages
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth.models import User
from django.core.paginator import Paginator
from django.db.models import Q
from django.shortcuts import get_object_or_404, redirect, render

from .models import Socio
from .models import UserProfile
import csv
from io import TextIOWrapper
from django.contrib.auth.models import User
from django.http import HttpResponse
import openpyxl
from openpyxl import Workbook


@login_required
@user_passes_test(lambda u: u.is_staff, login_url='/login/')
def listar_socios(request):
    q = request.GET.get('q', '').strip()
    estado = request.GET.get('estado', '').strip()

    socios = Socio.objects.select_related('user')
    if q:
        socios = socios.filter(
            Q(nombre__icontains=q)
            | Q(apellido__icontains=q)
            | Q(email__icontains=q)
            | Q(user__username__icontains=q)
        )
    if estado:
        socios = socios.filter(estado=estado)

    paginator = Paginator(socios.order_by('-fecha_ingreso'), 10)
    page_obj = paginator.get_page(request.GET.get('page'))

    # Asegurar que cada usuario listado tenga un UserProfile para evitar errores en plantillas
    for s in page_obj.object_list:
        try:
            UserProfile.objects.get_or_create(user=s.user)
        except Exception:
            pass

    return render(request, 'socios/socios.html', {
        'page_obj': page_obj,
        'q': q,
        'estado': estado,
        'is_admin': request.user.is_staff,
    })


@login_required
@user_passes_test(lambda u: u.is_staff, login_url='/login/')
def crear_socio(request):
    if request.method != 'POST':
        return redirect('socios:listar_socios')

    username = request.POST.get('username', '').strip()
    nombre = request.POST.get('nombre', '').strip()
    apellido = request.POST.get('apellido', '').strip()
    email = request.POST.get('email', '').strip()
    telefono = request.POST.get('telefono', '').strip()
    ciudad = request.POST.get('ciudad', '').strip()
    direccion = request.POST.get('direccion', '').strip()
    password = request.POST.get('password', '')

    if not username or not nombre or not apellido or not email or not password:
        messages.error(request, 'Completa los campos obligatorios.')
        return redirect('socios:listar_socios')

    if User.objects.filter(username=username).exists():
        messages.error(request, 'El nombre de usuario ya existe.')
        return redirect('socios:listar_socios')

    user = User.objects.create_user(username=username, email=email, password=password)
    user.first_name = nombre
    user.last_name = apellido
    user.save()
    Socio.objects.create(
        user=user,
        nombre=nombre,
        apellido=apellido,
        email=email,
        telefono=telefono,
        ciudad=ciudad,
        direccion=direccion,
    )
    messages.success(request, 'Socio registrado correctamente.')
    return redirect('socios:listar_socios')


@login_required
def perfil_socio(request):
    try:
        socio = request.user.socio_profile
    except Socio.DoesNotExist:
        messages.error(request, 'No se encontró el perfil de socio.')
        return redirect('/')

    entregas = socio.entregas_souvenir.select_related('entregado_por').all()
    paginator = Paginator(entregas, 10)
    page_obj = paginator.get_page(request.GET.get('page'))

    return render(request, 'socios/perfil.html', {
        'socio': socio,
        'page_obj': page_obj,
        'is_admin': request.user.is_staff,
    })


@login_required
def subir_foto(request):
    if request.method != 'POST':
        return redirect('socios:perfil_socio')

    foto = request.FILES.get('foto')
    user_id = request.POST.get('user_id')

    # Si el usuario es admin puede subir foto para otro usuario
    if user_id and request.user.is_staff:
        from django.contrib.auth.models import User
        target = User.objects.filter(id=user_id).first()
        if not target:
            messages.error(request, 'Usuario no encontrado.')
            return redirect('socios:perfil_socio')
        profile, _ = UserProfile.objects.get_or_create(user=target)
    else:
        profile, _ = UserProfile.objects.get_or_create(user=request.user)

    if foto:
        profile.foto = foto
        profile.save()
        messages.success(request, 'Foto de perfil actualizada.')
    else:
        messages.error(request, 'No se recibió archivo.')

    return redirect('socios:perfil_socio')


@login_required
@user_passes_test(lambda u: u.is_staff, login_url='/login/')
def crear_admin(request):
    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        email = request.POST.get('email', '').strip()
        password = request.POST.get('password', '')
        first_name = request.POST.get('first_name', '').strip()
        last_name = request.POST.get('last_name', '').strip()
        if not username or not email or not password:
            messages.error(request, 'Completa los campos obligatorios.')
            return redirect('socios:crear_admin')
        if User.objects.filter(username=username).exists():
            messages.error(request, 'El nombre de usuario ya existe.')
            return redirect('socios:crear_admin')
        user = User.objects.create_user(username=username, email=email, password=password)
        user.first_name = first_name
        user.last_name = last_name
        user.is_staff = True
        user.save()
        messages.success(request, 'Administrador creado correctamente.')
        return redirect('socios:listar_admins')
    return redirect('socios:listar_admins')


@login_required
@user_passes_test(lambda u: u.is_staff, login_url='/login/')
def importar_socios(request):
    if request.method == 'POST':
        f = request.FILES.get('file')
        if not f:
            messages.error(request, 'Sube un archivo CSV.')
            return redirect('socios:listar_socios')
        try:
            text = TextIOWrapper(f.file, encoding='utf-8')
            reader = csv.DictReader(text)
            created = 0
            for row in reader:
                username = row.get('username') or row.get('usuario') or ''
                nombre = row.get('nombre') or ''
                apellido = row.get('apellido') or ''
                email = row.get('email') or ''
                password = row.get('password') or User.objects.make_random_password()
                telefono = row.get('telefono') or ''
                ciudad = row.get('ciudad') or ''
                direccion = row.get('direccion') or ''
                if not username or User.objects.filter(username=username).exists():
                    continue
                user = User.objects.create_user(username=username, email=email, password=password)
                user.first_name = nombre
                user.last_name = apellido
                user.save()
                Socio.objects.create(user=user, nombre=nombre, apellido=apellido, email=email, telefono=telefono, ciudad=ciudad, direccion=direccion)
                created += 1
            messages.success(request, f'Socios importados: {created}')
        except Exception as e:
            messages.error(request, f'Error al procesar el archivo: {e}')
    return redirect('socios:listar_socios')


@login_required
@user_passes_test(lambda u: u.is_staff, login_url='/login/')
def descargar_plantilla_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = 'socios'
    headers = ['username', 'nombre', 'apellido', 'email', 'password', 'telefono', 'ciudad', 'direccion']
    ws.append(headers)
    # ejemplo fila
    ws.append(['jdoe', 'Juan', 'Doe', 'jdoe@example.com', 'Passw0rd!', '71234567', 'Oruro', 'Dirección 123'])
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=socios_plantilla.xlsx'
    wb.save(response)
    return response


@login_required
@user_passes_test(lambda u: u.is_staff, login_url='/login/')
def importar_socios_xlsx(request):
    if request.method == 'POST':
        f = request.FILES.get('file')
        if not f:
            messages.error(request, 'Sube un archivo .xlsx')
            return redirect('socios:listar_socios')
        try:
            wb = openpyxl.load_workbook(f)
            ws = wb.active
            created = 0
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0:
                    continue
                username, nombre, apellido, email, password, telefono, ciudad, direccion = [ (c or '') for c in row[:8] ]
                if not username or User.objects.filter(username=username).exists():
                    continue
                if not password:
                    password = User.objects.make_random_password()
                user = User.objects.create_user(username=username, email=email, password=password)
                user.first_name = nombre
                user.last_name = apellido
                user.save()
                Socio.objects.create(user=user, nombre=nombre, apellido=apellido, email=email, telefono=telefono, ciudad=ciudad, direccion=direccion)
                created += 1
            messages.success(request, f'Socios importados desde XLSX: {created}')
        except Exception as e:
            messages.error(request, f'Error al procesar xlsx: {e}')
    return redirect('socios:listar_socios')


@login_required
@user_passes_test(lambda u: u.is_staff, login_url='/login/')
def listar_admins(request):
    admins = User.objects.filter(is_staff=True).order_by('username')
    paginator = Paginator(admins, 20)
    page_obj = paginator.get_page(request.GET.get('page'))
    return render(request, 'admins/admins.html', {'page_obj': page_obj})


@login_required
@user_passes_test(lambda u: u.is_staff, login_url='/login/')
def ver_admin(request, user_id):
    return redirect('socios:listar_admins')


@login_required
@user_passes_test(lambda u: u.is_staff, login_url='/login/')
def editar_admin(request, user_id):
    user = get_object_or_404(User, id=user_id, is_staff=True)
    if request.method == 'POST':
        user.username = request.POST.get('username', user.username).strip()
        user.email = request.POST.get('email', user.email).strip()
        user.first_name = request.POST.get('first_name', user.first_name).strip()
        user.last_name = request.POST.get('last_name', user.last_name).strip()
        password = request.POST.get('password')
        if password:
            user.set_password(password)
        user.save()
        messages.success(request, 'Administrador actualizado.')
        return redirect('socios:listar_admins')
    return redirect('socios:listar_admins')


@login_required
@user_passes_test(lambda u: u.is_staff, login_url='/login/')
def eliminar_admin(request, user_id):
    user = get_object_or_404(User, id=user_id, is_staff=True)
    if request.method == 'POST':
        user.delete()
        messages.success(request, 'Administrador eliminado.')
        return redirect('socios:listar_admins')
    return redirect('socios:listar_admins')


@login_required
def mis_souvenirs(request):
    try:
        socio = request.user.socio_profile
    except Socio.DoesNotExist:
        messages.error(request, 'No se encontró perfil de socio.')
        return redirect('/')
    entregas = socio.entregas_souvenir.select_related('souvenir', 'entregado_por').all()
    paginator = Paginator(entregas, 10)
    page_obj = paginator.get_page(request.GET.get('page'))
    return render(request, 'socios/mis_souvenirs.html', {'page_obj': page_obj})


@login_required
@user_passes_test(lambda u: u.is_staff, login_url='/login/')
def editar_socio(request, socio_id):
    if request.method != 'POST':
        return redirect('socios:listar_socios')

    socio = get_object_or_404(Socio, id=socio_id)
    socio.nombre = request.POST.get('nombre', '').strip()
    socio.apellido = request.POST.get('apellido', '').strip()
    socio.email = request.POST.get('email', '').strip()
    socio.telefono = request.POST.get('telefono', '').strip()
    socio.ciudad = request.POST.get('ciudad', '').strip()
    socio.direccion = request.POST.get('direccion', '').strip()
    socio.observacion = request.POST.get('observacion', '').strip()
    socio.save()
    messages.success(request, 'Datos del socio actualizados.')
    return redirect('socios:listar_socios')


@login_required
@user_passes_test(lambda u: u.is_staff, login_url='/login/')
def activar_socio(request, socio_id):
    socio = get_object_or_404(Socio, id=socio_id)
    socio.estado = 'activo'
    socio.save()
    messages.success(request, 'Socio activado.')
    return redirect('socios:listar_socios')


@login_required
@user_passes_test(lambda u: u.is_staff, login_url='/login/')
def desactivar_socio(request, socio_id):
    socio = get_object_or_404(Socio, id=socio_id)
    socio.estado = 'inactivo'
    socio.save()
    messages.success(request, 'Socio desactivado.')
    return redirect('socios:listar_socios')


@login_required
@user_passes_test(lambda u: u.is_staff, login_url='/login/')
def eliminar_socio(request, socio_id):
    socio = get_object_or_404(Socio, id=socio_id)
    socio.user.delete()
    socio.delete()
    messages.success(request, 'Socio eliminado definitivamente.')
    return redirect('socios:listar_socios')
